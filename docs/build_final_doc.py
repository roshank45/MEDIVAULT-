import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def add_image_safely(doc, img_path, caption_text, width=Inches(6.2)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(71, 85, 105)

def build_final_report():
    doc = Document()

    # Margins: 1 inch
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    TEAL = RGBColor(13, 148, 136)
    DARK_NAVY = RGBColor(15, 23, 42)
    SLATE = RGBColor(71, 85, 105)

    MEDIA_DIR = r"C:\Users\boomi\.gemini\antigravity\brain\40c7053b-ff62-4798-bb18-662b0d834cb6\.user_uploaded"

    # Title
    t_para = doc.add_paragraph()
    t_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = t_para.add_run("MediVault")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(28)
    r_t.font.bold = True
    r_t.font.color.rgb = TEAL

    sub_para = doc.add_paragraph()
    sub_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_para.add_run("Centralized Patient Medical Records & Medication Safety System\nFinal Project Review & Comprehensive Module Execution Report")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(15)
    r_sub.font.bold = True
    r_sub.font.color.rgb = DARK_NAVY

    meta_para = doc.add_paragraph()
    meta_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = meta_para.add_run("Academic Review Document with Full Code, Execution Guide & Verified Live Terminal Outputs\nStack: Next.js 14 • FastAPI (Python 3.13) • Oracle Database 11g XE (3NF Normalized)")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(10.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = SLATE

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # =========================================================================
    # 1. INTRODUCTION
    # =========================================================================
    doc.add_heading("1. Introduction & Project Architecture", level=1).runs[0].font.color.rgb = DARK_NAVY

    p = doc.add_paragraph()
    p.add_run("1.1 Problem Statement: ").bold = True
    p.add_run(
        "In modern healthcare, medical records are isolated across disparate hospital silos. "
        "When patients visit new clinics or emergency rooms in unconscious or acute states, doctors have zero visibility into their historical diagnoses, "
        "vital trends, past treatments, or critical allergies. This fragmentation directly causes fatal adverse drug reactions (ADRs), "
        "duplicate therapies, and overlooked allergies (such as penicillin anaphylaxis)."
    )

    p = doc.add_paragraph()
    p.add_run("1.2 Solution: ").bold = True
    p.add_run(
        "MediVault is an enterprise-grade centralized medical records and clinical safety platform that provides authorized doctors with longitudinal "
        "patient history across participating hospitals. It issues unique patient identifiers (MV-YYYY-NNNNNN), performs automated real-time "
        "Clinical Decision Support (CDSS) for medication safety, and features a privacy-preserving Two-Tier Emergency QR Code profile with "
        "break-glass access controls."
    )

    # Tech Stack Table
    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Layer", "Technology Used", "Role in MediVault"]
    col_w = [Inches(1.2), Inches(2.0), Inches(3.3)]
    for i, h in enumerate(headers):
        c = table.cell(0, i)
        c.text = h
        set_cell_background(c, "0D9488")
        r = c.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)

    rows = [
        ("Frontend", "Next.js 14, React, TypeScript, Tailwind CSS", "Clinical dashboard, vitals charts, emergency QR view, break-glass modal."),
        ("Backend API", "FastAPI, Uvicorn, Pydantic v2, JWT, Passlib (Bcrypt)", "REST APIs, Role-Based Access Control (RBAC), medication safety algorithms, audit logging."),
        ("Database", "Oracle Database 11g Express Edition (XE), SQL, PL/SQL", "18 normalized tables (3NF), foreign keys, sequences, triggers, check constraints, views."),
        ("Database Driver", "pyodbc / Native 64-bit Oracle XE ODBC Driver", "Thread-safe pooled execution between Python async workers and Oracle on localhost:1521/XE.")
    ]
    for r_idx, r_data in enumerate(rows, start=1):
        for c_idx, val in enumerate(r_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            run = cell.paragraphs[0].runs[0]
            run.font.size = Pt(9)
            if c_idx == 0:
                run.font.bold = True

    for row in table.rows:
        for idx, w in enumerate(col_w):
            row.cells[idx].width = w

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================================
    # 2. MODULE-BY-MODULE EXPLANATION & LIVE OUTPUTS
    # =========================================================================
    doc.add_heading("2. Detailed Technical Modules, Code Locations & Live Execution Outputs", level=1).runs[0].font.color.rgb = DARK_NAVY

    # --- MODULE 1 ---
    doc.add_heading("Module 1: Patient Registration and Authentication", level=2).runs[0].font.color.rgb = TEAL
    p = doc.add_paragraph()
    p.add_run("• Functional Role: ").bold = True
    p.add_run("Manages identity boundaries, Unique Patient ID generation (MV-YYYY-NNNNNN), Doctor self-service onboarding, Admin provisioning, and RBAC security enforcement (ADMIN, DOCTOR, EMERGENCY_CLINICIAN, PATIENT).\n")
    p.add_run("• Where to Find Code: ").bold = True
    p.add_run("backend/app/api/auth.py, backend/app/api/patients.py, backend/app/core/security.py, database/01_create_schema.sql (Tables: APP_USER, PATIENT, HOSPITAL).\n")
    p.add_run("• Standalone Execution Command: ").bold = True
    p.add_run("python scripts/demo_module1.py")

    add_image_safely(
        doc,
        os.path.join(MEDIA_DIR, "media_1792732455582_b31eb836.png"),
        "Figure 1.1: Live Terminal Output for Module 1 (Part 1) — Doctor Registration (POST /api/auth/doctor-signup) & Patient ID Generation (MV-2026-000017)."
    )
    add_image_safely(
        doc,
        os.path.join(MEDIA_DIR, "media_1792732299534_ed2c2c99.png"),
        "Figure 1.2: Live Terminal Output for Module 1 (Part 2) — Bcrypt Password Login (200 OK) & RBAC Enforcement (Doctor restricted from Admin Audit: 403 FORBIDDEN)."
    )

    # --- MODULE 2 ---
    doc.add_heading("Module 2: Medical History Management", level=2).runs[0].font.color.rgb = TEAL
    p = doc.add_paragraph()
    p.add_run("• Functional Role: ").bold = True
    p.add_run("Manages multi-hospital consultations, structured ICD-10 diagnoses, longitudinal vitals tracking (Systolic/Diastolic BP, Glucose, SpO2), diagnostic test reports, and chronological timeline aggregation.\n")
    p.add_run("• Where to Find Code: ").bold = True
    p.add_run("backend/app/api/visits.py, backend/app/api/patients.py, backend/app/api/vitals.py, database/01_create_schema.sql (Tables: VISIT, DIAGNOSIS, VITAL_RECORD; View: VW_PATIENT_TIMELINE_SUMMARY).\n")
    p.add_run("• Standalone Execution Command: ").bold = True
    p.add_run("python scripts/demo_module2.py")

    add_image_safely(
        doc,
        os.path.join(MEDIA_DIR, "media_1792732299552_a03c57aa.png"),
        "Figure 2.1: Live Terminal Output for Module 2 (Part 1) — Patient Search & Consultation Recording at City General Hospital (Visit ID: 15, Systolic BP: 146 mmHg)."
    )
    add_image_safely(
        doc,
        os.path.join(MEDIA_DIR, "media_1792732299571_2ef3a996.png"),
        "Figure 2.2: Live Terminal Output for Module 2 (Part 2) — Cross-Hospital History Aggregation across City General & Sunrise Medical Centre (15 vital data points)."
    )

    # --- MODULE 3 ---
    doc.add_heading("Module 3: Emergency QR Medical Profile & Break-Glass Access", level=2).runs[0].font.color.rgb = TEAL
    p = doc.add_paragraph()
    p.add_run("• Functional Role: ").bold = True
    p.add_run("Two-Tier Emergency Framework: Tier 1 serves public essential triage data (Blood Group, Critical Allergies, Emergency Contact) via QR code without authentication, suppressing private clinic notes. Tier 2 provides authenticated ER clinicians with a 4-hour break-glass override with mandatory justification and audit logging.\n")
    p.add_run("• Where to Find Code: ").bold = True
    p.add_run("backend/app/api/emergency.py, frontend/app/emergency/qr/[token]/page.tsx, frontend/components/BreakGlassModal.tsx, database/01_create_schema.sql (Tables: EMERGENCY_ACCESS_TOKEN, PATIENT_ACCESS, AUDIT_LOG).\n")
    p.add_run("• Standalone Execution Command: ").bold = True
    p.add_run("python scripts/demo_module3.py")

    add_image_safely(
        doc,
        os.path.join(MEDIA_DIR, "media_1792732299580_786f09d7.png"),
        "Figure 3.1: Live Terminal Output for Module 3 (Part 1) — Public Unauthenticated QR Code Query (Triage Data: O_POSITIVE, Penicillin Allergy, Private notes filtered)."
    )
    add_image_safely(
        doc,
        os.path.join(MEDIA_DIR, "media_1792732299567_b8ada708.png"),
        "Figure 3.2: Live Terminal Output for Module 3 (Part 2) — Break-Glass Emergency Override (4-Hour Window Granted) & Oracle AUDIT_LOG Event Registration."
    )

    # --- MODULE 4 ---
    doc.add_heading("Module 4: Medication Safety and Reporting", level=2).runs[0].font.color.rgb = TEAL
    p = doc.add_paragraph()
    p.add_run("• Functional Role: ").bold = True
    p.add_run("Real-time Clinical Decision Support System (CDSS) with 3 safety algorithms: (1) Allergy Conflict Engine (Penicillin allergy vs Amoxicillin), (2) Cross-Hospital Duplicate Therapy Engine (active Metformin overlap), (3) Drug-Drug Interaction Engine (Warfarin + Aspirin bleeding hazard), and regulatory immutable audit logging.\n")
    p.add_run("• Where to Find Code: ").bold = True
    p.add_run("backend/app/services/safety_service.py, backend/app/api/safety.py, backend/app/api/audit.py, database/01_create_schema.sql (Tables: ALLERGY, PRESCRIPTION, DRUG_INTERACTION_RULE, AUDIT_LOG).\n")
    p.add_run("• Standalone Execution Command: ").bold = True
    p.add_run("python scripts/demo_module4.py")

    add_image_safely(
        doc,
        os.path.join(MEDIA_DIR, "media_1792732590006_d956a880.png"),
        "Figure 4.1: Live Terminal Output for Module 4 (Part 1) — Safety Interceptions: ALLERGY_CONFLICT (Amoxil vs Penicillin), DUPLICATE_THERAPY (Metformin), DRUG_INTERACTION (Warfarin+Aspirin)."
    )
    add_image_safely(
        doc,
        os.path.join(MEDIA_DIR, "media_1792732590081_49a1b25c.png"),
        "Figure 4.2: Live Terminal Output for Module 4 (Part 2) — Regulatory Audit Trail in Oracle AUDIT_LOG Table showing Break-Glass and Authentication events."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================================
    # 3. HOW TO EXECUTE EACH MODULE SEPARATELY
    # =========================================================================
    doc.add_heading("3. Execution Guide for Academic Project Viva / Review", level=1).runs[0].font.color.rgb = DARK_NAVY
    p = doc.add_paragraph()
    p.add_run(
        "Each module can be executed separately in the terminal or demonstrated via the web user interface. "
        "Here are the exact commands to demonstrate each module individually:"
    )

    p_cmd = doc.add_paragraph()
    p_cmd.paragraph_format.left_indent = Inches(0.2)
    r_cmd = p_cmd.add_run(
        '# Step 1: Navigate to the project root in PowerShell\n'
        'cd C:\\Users\\boomi\\.gemini\\antigravity\\scratch\\medivault\n\n'
        '# Step 2: Execute Any Desired Module Independently\n'
        'python scripts/demo_module1.py   # Runs Module 1 (Registration & Auth)\n'
        'python scripts/demo_module2.py   # Runs Module 2 (Consultations & Vitals)\n'
        'python scripts/demo_module3.py   # Runs Module 3 (Emergency QR & Break-Glass)\n'
        'python scripts/demo_module4.py   # Runs Module 4 (Medication Safety & Audit)\n\n'
        '# Step 3: Run All 4 Modules Sequentially in One Viva Command\n'
        'python scripts/run_all_modules.py\n\n'
        '# Step 4: Run Automated Pytest Test Suite (19/19 Passing)\n'
        'pytest tests/test_backend.py -v'
    )
    r_cmd.font.name = "Consolas"
    r_cmd.font.size = Pt(8.5)
    r_cmd.font.color.rgb = DARK_NAVY

    # =========================================================================
    # 4. WEB APPLICATION USER INTERFACE OUTPUTS
    # =========================================================================
    doc.add_heading("4. Graphical User Interface (GUI) Screenshots", level=1).runs[0].font.color.rgb = DARK_NAVY

    add_image_safely(
        doc,
        os.path.join(MEDIA_DIR, "media_1792687998597_bbc97460.png"),
        "Figure 5.1: MediVault Clinical Dashboard — Harold Finch Profile (MV-2026-000001, Blood: O_POSITIVE) showing active cross-hospital prescriptions (Glucophage, Prinivil, Lipitor) and Documented Allergies Alert."
    )

    # Route Table
    rt_table = doc.add_table(rows=6, cols=3)
    rt_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rt_table.autofit = False

    r_headers = ["Module & Interface", "Local URL", "What to Demonstrate"]
    r_col_w = [Inches(1.8), Inches(2.3), Inches(2.4)]
    for i, h in enumerate(r_headers):
        c = rt_table.cell(0, i)
        c.text = h
        set_cell_background(c, "0D9488")
        r = c.paragraphs[0].runs[0]
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    r_rows = [
        ("Module 1: Doctor Sign Up", "http://localhost:3000/signup", "Doctor registration with medical license validation, specialty, and hospital selector."),
        ("Module 1: Admin Sign Up", "http://localhost:3000/admin/signup", "Administrator registration using verified master passkey (MEDIVAULT_ADMIN_2026)."),
        ("Module 2: Longitudinal Record", "http://localhost:3000/patients/1", "Multi-hospital chronological timeline, vitals trend chart, and consultation recording."),
        ("Module 3: Emergency QR Card", "http://localhost:3000/emergency/qr/demo-...", "Public triage view with blood group and allergy alerts; private notes suppressed."),
        ("Module 4: Oracle Audit Explorer", "http://localhost:3000/admin/audit", "Live Oracle audit trail table displaying logins, queries, and break-glass overrides.")
    ]
    for r_idx, r_data in enumerate(r_rows, start=1):
        for c_idx, val in enumerate(r_data):
            cell = rt_table.cell(r_idx, c_idx)
            cell.text = val
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            run = cell.paragraphs[0].runs[0]
            run.font.size = Pt(9)
            if c_idx == 0:
                run.font.bold = True

    for row in rt_table.rows:
        for idx, w in enumerate(r_col_w):
            row.cells[idx].width = w

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # =========================================================================
    # 5. CONCLUSION
    # =========================================================================
    doc.add_heading("5. Conclusion & Project Accomplishments", level=1).runs[0].font.color.rgb = DARK_NAVY
    p = doc.add_paragraph()
    p.add_run(
        "The MediVault project successfully delivers a complete, production-modeled centralized medical records and medication safety platform. "
        "Unlike basic frontend prototypes, MediVault features real end-to-end integration across a Next.js 14 frontend, FastAPI Python backend, "
        "and Oracle Database 11g Express Edition.\n\n"
        "Key Highlights & Academic Milestones Achieved:\n"
        "1. True 3NF Relational Database Normalization: 18 normalized tables with foreign keys, sequences, triggers, check constraints, and views executed live in Oracle Database 11.2 XE.\n"
        "2. Clinical Safety Guardrails: Automated real-time Clinical Decision Support algorithms preventing allergy conflicts and cross-hospital duplicate prescriptions.\n"
        "3. Two-Tier Emergency Architecture: A tamper-resistant QR-based triage profile coupled with an authenticated break-glass emergency workflow for unconscious patients.\n"
        "4. Full Regulatory Auditability: Immutable logging of sensitive clinical queries, break-glass events, and administrative changes.\n"
        "5. Automated Test Rigor: 19 out of 19 automated tests passing with 100% success rate across authentication, RBAC, clinical history, emergency overrides, and safety checks.\n\n"
        "Disclaimer: MediVault is an academic prototype developed for educational demonstration. Production deployment requires full HIPAA/GDPR certification, FHIR/HL7 interoperability standards, and hardware-backed key management."
    )

    out_file = r"C:\Users\boomi\.gemini\antigravity\scratch\medivault\docs\MediVault_Final_Project_Report_With_All_Outputs.docx"
    doc.save(out_file)
    print(f"SUCCESS: Generated comprehensive report at {out_file}")

if __name__ == "__main__":
    build_final_report()
